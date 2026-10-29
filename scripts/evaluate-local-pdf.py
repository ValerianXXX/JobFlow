"""Compare installed converters on the same six synthetic DOCX files.

No model calls, applicant files, external requests, or MATERIALS_VALIDATED writes.
"""
from __future__ import annotations

import argparse
import base64
import hashlib
import json
import os
from pathlib import Path
import platform
import re
import sys
import time

PROJECT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT / "src"))

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
import pdfplumber

from jobops.document_builder import export_docx_to_pdf
from jobops.errors import JobOpsError
from jobops.source_review import source_page
from jobops.util import iso_utc, sha256_bytes


CASES = [
    {"id": "english", "language": "en", "layout": "paragraphs", "expected_pages": 1,
     "facts": ["In 2024, I checked 12 stores. The team saved $35,000; I verified the data, not the savings.",
               "From 2023-06 to 2024-02, I reviewed 48 reports; the reported error rate fell from 2.5% to 1.5%."]},
    {"id": "chinese", "language": "zh", "layout": "paragraphs", "expected_pages": 1,
     "facts": ["2024 年核验 12 家门店的数据。团队节约 3.5 万元，个人仅负责数据核验。",
               "2023 年 6 月至 2024 年 2 月，审核 48 份报告；记录的错误率从 2.5% 降至 1.5%。"]},
    {"id": "mixed", "language": "mixed", "layout": "paragraphs", "expected_pages": 1,
     "facts": ["2024 年完成 12-store audit；team savings 为 $35,000，个人仅负责核验。",
               "Project Alpha: 48 reports，2023-06 至 2024-02；error rate 2.5% → 1.5%。"]},
    {"id": "table", "language": "mixed", "layout": "table", "expected_pages": 1,
     "facts": ["Alpha / 2022-01–2022-12 / Analyst / 18 reports", "Beta / 2023-03–2024-02 / 实习生 / 24 reports",
               "The 20% result belongs to the Beta team; my role was data checking."]},
    {"id": "columns", "language": "en", "layout": "columns", "expected_pages": 1,
     "facts": ["Alpha: I tested 12 cases in 2024; the team reported 20% fewer errors.",
               "Beta: I documented 24 cases in 2025; no financial outcome was measured."]},
    {"id": "two-pages", "language": "mixed", "layout": "page_break", "expected_pages": 2,
     "facts": ["Project Alpha / 2024 / 12 stores / $35,000 team savings / verification only.",
               "项目 Beta / 2025 年 / 24 份报告 / 未测量收入 / 个人负责记录。"]},
]


def build_case(case: dict, target: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.7)
    style = document.styles["Normal"]
    style.font.name, style.font.size = "Arial", Pt(11)
    style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    document.core_properties.author = document.core_properties.last_modified_by = ""
    document.add_heading("SYNTHETIC CONVERSION FIXTURE", 0)
    document.add_paragraph("Not an applicant resume. All facts below are invented test data.")
    if case["layout"] in {"table", "columns"}:
        count = 2 if case["layout"] == "columns" else 1
        table = document.add_table(rows=1 if count == 2 else len(case["facts"]), cols=count)
        table.style = "Table Grid"
        for index, fact in enumerate(case["facts"]):
            table.cell(0 if count == 2 else index, index if count == 2 else 0).text = fact
    else:
        for index, fact in enumerate(case["facts"]):
            if index and case["layout"] == "page_break":
                document.add_page_break()
            document.add_paragraph(fact)
    footer = section.footer.paragraphs[0]
    footer.add_run("Synthetic fixture - page ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    document.save(target)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--providers", nargs="+", choices=("word", "libreoffice"), required=True)
    parser.add_argument("--output", type=Path, default=PROJECT / "reports" / "local-pdf-comparison")
    args = parser.parse_args()
    args.output.mkdir(parents=True, exist_ok=True)
    results = []
    script = PROJECT / ".agents" / "skills" / "job-application-operator" / "scripts" / "export-docx-pdf.ps1"
    for case in CASES:
        source = args.output / f"{case['id']}.docx"
        build_case(case, source)
        for provider in dict.fromkeys(args.providers):
            os.environ["JOBFLOW_PDF_EXPORT_PROVIDER"] = provider
            target = args.output / provider / f"{case['id']}.pdf"
            row = {"case": case["id"], "language": case["language"], "layout": case["layout"], "provider": provider,
                   "input_sha256": sha256_bytes(source.read_bytes()), "expected_pages": case["expected_pages"]}
            start = time.monotonic()
            try:
                export_docx_to_pdf(source, target, script)
                row["conversion_ms"] = round((time.monotonic() - start) * 1000)
                raw = target.read_bytes()
                with pdfplumber.open(target) as pdf:
                    texts = [page.extract_text() or "" for page in pdf.pages]
                normalized = re.sub(r"\s+", "", "\n".join(texts))
                row["missing_facts"] = [index + 1 for index, fact in enumerate(case["facts"])
                                        if re.sub(r"\s+", "", fact) not in normalized]
                row["blank_text_pages"] = [index + 1 for index, text in enumerate(texts) if not text.strip()]
                row["pages"], row["pdf_sha256"] = len(texts), sha256_bytes(raw)
                row["render_sha256"] = []
                for page in range(1, len(texts) + 1):
                    rendered = source_page(raw, ".pdf", expected_hash=row["pdf_sha256"], page=page)
                    image = base64.b64decode(rendered["image_data_url"].split(",", 1)[1])
                    target.with_name(f"{case['id']}-{page}.png").write_bytes(image)
                    row["render_sha256"].append(sha256_bytes(image))
                row["checks_passed"] = not row["missing_facts"] and not row["blank_text_pages"] and len(texts) == case["expected_pages"]
            except JobOpsError as exc:
                row.update(checks_passed=False, error_code=exc.code, failure_stage=exc.details.get("failure_stage"),
                           conversion_ms=round((time.monotonic() - start) * 1000))
            results.append(row)
            print(json.dumps(row), flush=True)
    receipt = {"schema_version": 1, "observed_at": iso_utc(), "python": platform.python_version(),
               "platform": platform.system(), "synthetic_only": True, "model_calls": 0,
               "model_call_cost_usd": 0, "software_licensing_cost_measured": False,
               "fixture_definition_sha256": "sha256:" + hashlib.sha256(json.dumps(CASES, ensure_ascii=False, sort_keys=True).encode()).hexdigest(),
               "visual_review": "PENDING_SEPARATE_REVIEW", "materials_validated": False,
               "results": results}
    (args.output / "result.json").write_text(json.dumps(receipt, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return 0 if all(item["checks_passed"] for item in results) else 1


if __name__ == "__main__":
    raise SystemExit(main())
