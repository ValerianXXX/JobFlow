"""Author 36 original synthetic inputs; run with the document authoring runtime.

reportlab is an authoring dependency only. Tests consume frozen, hash-checked
files and never regenerate them or install a new product dependency.
"""
from __future__ import annotations

import argparse
import hashlib
import io
from itertools import groupby
import json
from pathlib import Path

from docx import Document
from PIL import Image, ImageDraw, ImageFont, ImageFilter
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfgen import canvas

KINDS = ("single", "columns", "staggered_columns", "wrapped", "ruled_table", "two_pages",
         "docx", "same_organization_roles", "scan", "scanned_columns", "portfolio", "weak_image")
LANGUAGES = ("en", "zh", "mixed")


def entities(language: str, index: int) -> list[dict]:
    n = index + 1
    if language == "en":
        values = [
            (f"Synthetic Cedar {n:02d}", "Analyst | 2023-2024", [f"Team reviewed {12+n} stores.", f"Applicant prepared {6+n} reports."], [str(12+n), str(6+n)]),
            (f"Synthetic Harbor {n:02d}", "Researcher | 2021-2022", [f"Team managed a ${12500+n} budget.", f"Applicant tested {24+n} scenarios."], [str(12500+n), str(24+n)])]
    elif language == "zh":
        values = [
            (f"合成青杉实验室{n:02d}", "分析员 | 2023-2024", [f"团队核验了{12+n}家门店。", f"申请人编写了{6+n}份报告。"], [str(12+n), str(6+n)]),
            (f"合成海港研究所{n:02d}", "研究员 | 2021-2022", [f"团队管理预算{12500+n}美元。", f"申请人测试了{24+n}个场景。"], [str(12500+n), str(24+n)])]
    else:
        values = [
            (f"Synthetic Cedar 合成{n:02d}", "分析员 Analyst | 2023-2024", [f"团队 reviewed {12+n} 家门店。", f"申请人 prepared {6+n} reports."], [str(12+n), str(6+n)]),
            (f"Synthetic Harbor 合成{n:02d}", "研究员 Researcher | 2021-2022", [f"团队 managed ${12500+n} budget.", f"申请人 tested {24+n} scenarios."], [str(12500+n), str(24+n)])]
    return [{"organization": org, "role_dates": dates, "years": ["2023", "2024"] if i == 0 else ["2021", "2022"],
             "lines": [org, dates, *facts],
             "facts": [{"statement": fact, "number": number, "owner": "team" if j == 0 else "applicant"}
                       for j, (fact, number) in enumerate(zip(facts, numbers))]}
            for i, (org, dates, facts, numbers) in enumerate(values)]


def positions(kind: str, groups: list[dict]) -> list[tuple[int, float, float, str, int, int]]:
    result = []
    for entity, group in enumerate(groups):
        for line, text in enumerate(group["lines"]):
            page, x, y = 1, 44, 680 - entity * 150 - line * 28
            if kind in {"columns", "staggered_columns", "scanned_columns"}:
                x, y = 38 + entity * 300, 680 - line * 32
                if kind == "staggered_columns":
                    y -= entity * 10
            elif kind == "two_pages":
                page, y = entity + 1, 680 - line * 28
            elif kind == "ruled_table":
                x, y = 38 + entity * 300, 680 - line * 32
            if kind == "wrapped" and line == 2:
                # Exact same fact, deliberately divided over physical lines.
                cut = max(1, text.rfind(" ", 0, len(text) // 2 + 1)) if " " in text else len(text) // 2
                result.append((page, x, y, text[:cut], entity, line))
                result.append((page, x, y - 13, text[cut:].lstrip(), entity, line))
            else:
                result.append((page, x, y, text, entity, line))
    # Content stream follows rows, not entity order, to exercise real layout.
    return sorted(result, key=lambda row: (row[0], -row[2], row[1]))


def text_pdf(kind: str, groups: list[dict]) -> bytes:
    out = io.BytesIO()
    c = canvas.Canvas(out, pagesize=(612, 792), invariant=1, pageCompression=1)
    c.setAuthor("JobFlow synthetic fixture generator")
    c.setTitle("Synthetic material - not a real applicant")
    rows = positions(kind, groups)
    for page in range(1, max(row[0] for row in rows) + 1):
        c.setFont("Helvetica", 9)
        c.drawString(38, 750, "SYNTHETIC MATERIAL - NOT A REAL APPLICANT")
        c.drawString(38, 35, f"Original synthetic fixture | Page {page}")
        for number, x, y, text, _, _ in rows:
            if number == page:
                runs = [("Helvetica" if ascii_run else "STSong-Light", "".join(chars))
                        for ascii_run, chars in groupby(text, key=lambda char: ord(char) < 128)]
                if sum(pdfmetrics.stringWidth(run, font, 10) for font, run in runs) > (270 if "columns" in kind or kind == "ruled_table" else 530):
                    raise ValueError("Synthetic line exceeds its intended region")
                for font, run in runs:
                    c.setFont(font, 10)
                    c.drawString(x, y, run)
                    x += pdfmetrics.stringWidth(run, font, 10)
        if kind == "ruled_table":
            for y in (695, 663, 631, 599, 567):
                c.line(30, y, 585, y)
            for x in (30, 324, 585):
                c.line(x, 567, x, 695)
        c.showPage()
    c.save()
    return out.getvalue()


def raster(kind: str, groups: list[dict], font_path: Path) -> Image.Image:
    image = Image.new("RGB", (1224, 1584), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype(str(font_path), 22)
    heading = ImageFont.truetype(str(font_path), 20)
    draw.text((76, 70), "SYNTHETIC MATERIAL - NOT A REAL APPLICANT", font=heading, fill="#163248")
    for _, x, y, text, _, _ in positions(kind, groups):
        draw.text((int(x * 2), int((792-y) * 2)), text, font=font,
                  fill="#777777" if kind == "weak_image" else "#111111")
    if kind == "weak_image":
        image = image.filter(ImageFilter.GaussianBlur(.7))
    return image


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--raster-font", type=Path, required=True)
    args = parser.parse_args()
    args.output.mkdir(parents=True, exist_ok=False)
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    cases = []
    for language_index, language in enumerate(LANGUAGES):
        for index, kind in enumerate(KINDS):
            groups = entities(language, index)
            case_id = f"{language}-{kind}"
            same_organization = kind == "same_organization_roles"
            if same_organization:
                groups[1]["organization"] = groups[0]["organization"]
                groups[1]["lines"][0] = groups[0]["organization"]
            extension = ".docx" if kind in {"docx", "same_organization_roles"} else ".png" if kind in {"portfolio", "weak_image"} else ".pdf"
            if extension == ".docx":
                doc = Document()
                doc.add_heading("SYNTHETIC MATERIAL - NOT A REAL APPLICANT", 0)
                for group in groups:
                    for line in group["lines"]:
                        doc.add_paragraph(line)
                out = io.BytesIO()
                doc.save(out)
                data = out.getvalue()
            elif kind in {"scan", "scanned_columns", "portfolio", "weak_image"}:
                out = io.BytesIO()
                raster(kind, groups, args.raster_font).save(out, format="PDF" if extension == ".pdf" else "PNG", resolution=144)
                data = out.getvalue()
            else:
                data = text_pdf(kind, groups)
            filename = case_id + extension
            (args.output / filename).write_bytes(data)
            cases.append({"id": case_id, "language": language, "layout": kind,
                          "split": "held_out" if (index + language_index) % 3 == 2 else "development",
                          "file": filename, "sha256": hashlib.sha256(data).hexdigest(),
                          "page_count": 2 if kind == "two_pages" else None if extension == ".docx" else 1,
                          "entities": groups, "conflicting_dates_require_review": False,
                          "same_organization_different_roles": same_organization,
                          "text_only_expected": "OCR_OR_VISION_REQUIRED" if kind in {"scan", "scanned_columns", "portfolio", "weak_image"} else "READABLE"})
    manifest = {"schema_version": 1, "synthetic_only": True, "purpose": "Short layout probes; not the complete resume/portfolio acceptance corpus or a Claim accuracy benchmark",
                "split_rule": "(layout_index + language_index) % 3 == 2; 24 development / 12 held-out",
                "raster_font": {"name": args.raster_font.name, "sha256": hashlib.sha256(args.raster_font.read_bytes()).hexdigest()},
                "pdf_font": "Helvetica ASCII runs plus standard CID STSong-Light; renderer supplies a CJK font; neither font binary is distributed",
                "cases": cases}
    (args.output / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"cases": len(cases), "pdf": sum(c["file"].endswith(".pdf") for c in cases), "synthetic_only": True}))


if __name__ == "__main__":
    main()
