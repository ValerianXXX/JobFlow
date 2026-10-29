from __future__ import annotations

from copy import deepcopy
import io
import json
import os
import textwrap
import unittest
from unittest.mock import patch

import _support
from _support import PROJECT, project_temp
import test_onboarding_center as fixtures
from docx import Document
from jobops.document_builder import build_standard_resume
from jobops.document_capabilities import document_capabilities
from jobops.errors import JobOpsError
from jobops.external_claims import build_external_claim_set, claim_review_hash
from jobops.util import canonical_json, sha256_bytes
from workbench_fixtures import text_pdf


def approved_pdf_service(root):
    service, onboarding, store, _, _ = fixtures.OnboardingCenterTests().make_service(root)
    lines = ["SYNTHETIC RESUME - NOT A REAL APPLICANT"]
    for item in service._claim_bundle()["claims"]:
        lines.extend(textwrap.wrap(item["resume_statement"], width=80))
    raw = text_pdf([(48, 740 - index * 20, line) for index, line in enumerate(lines)])
    original = onboarding.import_bytes("master_resume_pdf", raw, synthetic=True)
    fixtures.OnboardingCenterTests.prepare_completion(service)
    service.complete(user_confirmed=True)
    approval = service.bootstrap()["external_claim_approval"]
    service.approve_external_claims({"user_confirmed": True, "expected_review_hash": approval["review_hash"],
                                     "allowed_uses": approval["allowed_uses"]})
    return service, onboarding, store, original


def generate(service, locale="en"):
    return service.create_standard_resume({"expected_review_hash": service.bootstrap()["standard_resume"]["review_hash"], "locale": locale})


class StandardResumeTests(unittest.TestCase):
    def test_pdf_to_docx_revision_preserves_original_evidence_and_requires_new_approvals(self):
        with project_temp() as root:
            service, onboarding, _, original = approved_pdf_service(root)
            old_ref, old_state = service.ensure_state()
            original_state_bytes = onboarding.read_bytes(old_ref)
            original_pdf_bytes = onboarding.read_bytes(str(original["secure_ref"]))
            result = generate(service)
            self.assertEqual(result["preview"]["model_calls"], 0)
            self.assertEqual(result["preview"]["cost_usd"], 0)
            self.assertFalse(result["materials_validated"])
            self.assertFalse(result["preview"]["preserves_original_layout"])
            self.assertEqual(generate(service)["draft_ref"], result["draft_ref"])
            document = Document(io.BytesIO(service.standard_resume_file(result)))
            paragraphs = [item.text for item in document.paragraphs]
            for section in result["preview"]["sections"]:
                for claim in section["items"]:
                    self.assertIn(claim["text"], paragraphs)
            self.assertEqual(onboarding.read_bytes(old_ref), original_state_bytes)
            adopt = {"draft_ref": result["draft_ref"], "expected_docx_sha256": result["docx_sha256"], "user_confirmed": True}
            service.adopt_standard_resume(adopt)
            new_ref, new_state = service.ensure_state()
            self.assertNotEqual(new_ref, old_ref)
            self.assertEqual(new_state["revision_number"], old_state["revision_number"] + 1)
            self.assertTrue(new_state["master_resume"]["editable_docx"])
            self.assertEqual(new_state["base_claim_source"]["secure_ref"], original["secure_ref"])
            self.assertEqual(onboarding.read_bytes(old_ref), original_state_bytes)
            self.assertEqual(onboarding.read_bytes(str(original["secure_ref"])), original_pdf_bytes)
            self.assertFalse(service.bootstrap()["external_claim_approval"]["current"])
            self.assertFalse(service.bootstrap()["tailoring_manifest"]["current"])
            self.assertFalse(service.adopt_standard_resume(adopt)["changed"])
            fixtures.OnboardingCenterTests.prepare_completion(service)
            service.complete(user_confirmed=True)
            approval = service.bootstrap()["external_claim_approval"]
            current = service.approve_external_claims({"user_confirmed": True, "expected_review_hash": approval["review_hash"], "allowed_uses": approval["allowed_uses"]})
            claims = json.loads(onboarding.read_bytes(current["claim_set_ref"]))
            self.assertTrue(all(binding["secure_ref"] == original["secure_ref"] for claim in claims["claims"] for binding in claim["source_bindings"]))
            proposal = service.tailoring_manifest_proposal()
            self.assertGreater(proposal["candidate_count"], 0)
            service.approve_tailoring_manifest({"user_confirmed": True, "expected_proposal_hash": proposal["proposal_hash"],
                "selections": [{"block_ref": item["block_ref"], "category": item["allowed_categories"][0]} for item in proposal["candidates"]]})
            self.assertEqual(service.bootstrap()["application_readiness"]["status"], "READY_FOR_OFFLINE_APPLICATION_PREPARATION")
            self.assertEqual(service.redacted_status()["real_external_actions"], 0)

    def test_missing_or_wrong_use_approval_cannot_generate(self):
        with project_temp() as root:
            service, onboarding, _, _ = approved_pdf_service(root)
            current = service.bootstrap()["external_claim_approval"]
            service.approve_external_claims({"user_confirmed": True, "expected_review_hash": current["review_hash"], "allowed_uses": ["cover_letter"]})
            self.assertFalse(service.bootstrap()["standard_resume"]["available"])
            with self.assertRaises(JobOpsError) as caught:
                generate(service)
            self.assertEqual(caught.exception.code, "EXTERNAL_CLAIM_APPROVAL_REQUIRED")
            self.assertIsNone(service._latest_ref("standard_resume_draft"))

    def test_stale_confirmation_and_foreign_reference_do_not_change_completed_profile(self):
        with project_temp() as root:
            service, onboarding, _, original = approved_pdf_service(root)
            old_ref, _ = service.ensure_state()
            before = onboarding.read_bytes(old_ref)
            result = generate(service)
            for changes, code in [({"user_confirmed": False}, "STANDARD_RESUME_CONFIRMATION_REQUIRED"),
                                  ({"expected_docx_sha256": "sha256:" + "0" * 64}, "STANDARD_RESUME_DRAFT_STALE"),
                                  ({"draft_ref": original["secure_ref"]}, "STANDARD_RESUME_DRAFT_INVALID")]:
                payload = {"draft_ref": result["draft_ref"], "expected_docx_sha256": result["docx_sha256"], "user_confirmed": True, **changes}
                with self.assertRaises(JobOpsError) as caught:
                    service.adopt_standard_resume(payload)
                self.assertEqual(caught.exception.code, code)
                self.assertEqual(onboarding.read_bytes(old_ref), before)
            service.start_revision()
            with self.assertRaises(JobOpsError):
                service.standard_resume_file(result)

    def test_revision_write_failure_keeps_completed_state_and_allows_retry(self):
        with project_temp() as root:
            service, onboarding, _, _ = approved_pdf_service(root)
            old_ref, _ = service.ensure_state()
            old_bytes = onboarding.read_bytes(old_ref)
            result = generate(service)
            payload = {**result, "expected_docx_sha256": result["docx_sha256"], "user_confirmed": True}
            with patch.object(service, "_write_index", side_effect=OSError("synthetic failure")):
                with self.assertRaises(JobOpsError) as caught:
                    service.adopt_standard_resume(payload)
            self.assertEqual(caught.exception.code, "ONBOARDING_REVISION_WRITE_FAILED")
            self.assertEqual(service.ensure_state()[0], old_ref)
            self.assertEqual(onboarding.read_bytes(old_ref), old_bytes)
            self.assertEqual(service.adopt_standard_resume(payload)["status"], "STANDARD_RESUME_ADOPTED")

    def test_corrupted_evidence_or_changed_profile_blocks_draft_reuse(self):
        with project_temp() as root:
            service, onboarding, _, original = approved_pdf_service(root)
            result = generate(service)
            profile_ref = service._latest_ref("candidate_profile")
            profile = json.loads(onboarding.read_bytes(profile_ref))
            profile["candidate_display_name"] = "Changed Synthetic Candidate"
            onboarding.rotate(profile_ref, canonical_json(profile))
            with self.assertRaises(JobOpsError) as caught:
                service.standard_resume_file(result)
            self.assertEqual(caught.exception.code, "STANDARD_RESUME_DRAFT_STALE")
            onboarding.rotate(str(original["secure_ref"]), b"changed evidence")
            with self.assertRaises(JobOpsError):
                generate(service)

    @patch.dict(os.environ, {"JOBFLOW_PDF_EXPORT_PROVIDER": "word"})
    def test_pdf_failure_keeps_docx_and_redacts_converter_errors(self):
        with project_temp() as root:
            service, _, _, _ = approved_pdf_service(root)
            result = generate(service)
            with patch("jobops.document_capabilities.word_registered", return_value=False):
                with self.assertRaises(JobOpsError) as caught:
                    service.standard_resume_pdf(result)
            self.assertEqual(caught.exception.code, "DOCUMENT_PDF_EXPORT_UNAVAILABLE")
            with patch("jobops.document_capabilities.word_registered", return_value=True), patch(
                "jobops.document_builder.export_docx_to_pdf", side_effect=JobOpsError("DOCX_PDF_EXPORT_FAILED", "sensitive document path")
            ):
                with self.assertRaises(JobOpsError) as caught:
                    service.standard_resume_pdf(result)
            self.assertEqual(caught.exception.code, "DOCUMENT_PDF_EXPORT_FAILED")
            self.assertNotIn("sensitive", json.dumps(caught.exception.as_dict()))
            self.assertEqual(sha256_bytes(service.standard_resume_file(result)), result["docx_sha256"])

    @patch.dict(os.environ, {"JOBFLOW_PDF_EXPORT_PROVIDER": "word"})
    def test_pdf_commit_failure_rolls_back_new_encrypted_output(self):
        with project_temp() as root:
            service, onboarding, _, _ = approved_pdf_service(root)
            result = generate(service)
            before = onboarding.read_bytes(result["draft_ref"])
            pdf = text_pdf([(48, 740, "SYNTHETIC CONVERTER CONTRACT")])
            def export(_docx, target, _script):
                target.write_bytes(pdf)
            with patch("jobops.document_capabilities.word_registered", return_value=True), patch(
                "jobops.document_builder.export_docx_to_pdf", side_effect=export
            ), patch.object(onboarding, "rotate", side_effect=OSError("synthetic commit failure")):
                with self.assertRaises(JobOpsError) as caught:
                    service.standard_resume_pdf(result)
            self.assertEqual(caught.exception.details["failure_stage"], "COMMIT")
            self.assertEqual(onboarding.read_bytes(result["draft_ref"]), before)
            self.assertIsNone(service._latest_ref("generated_resume_pdf"))
            self.assertEqual(sha256_bytes(service.standard_resume_file(result)), result["docx_sha256"])

    def test_bilingual_layout_preserves_exact_numbers_and_has_no_invented_sections(self):
        with project_temp() as root:
            evidence_hash = "sha256:" + "1" * 64
            wording = "2024 年负责 12 家门店的数据核验，团队节约 3.5 万元；个人仅负责核验。"
            claim = {"claim_id": "CLM-SYNTHETIC-ZH", "category": "project", "claim_kind": "responsibility", "decision": "CONFIRMED",
                     "statement": wording, "source_bindings": [{"kind": "MASTER_RESUME", "secure_ref": "secure-ref:SYNTHETIC_MASTER", "content_sha256": evidence_hash}]}
            approved = build_external_claim_set(onboarding_state_ref="secure-ref:SYNTHETIC_STATE", profile_ref="secure-ref:SYNTHETIC_PROFILE",
                master_resume={"secure_ref": "secure-ref:SYNTHETIC_MASTER", "sha256": evidence_hash, "editable_docx": False},
                claims=[claim], allowed_uses=["resume"], expected_review_hash=claim_review_hash([claim], evidence_hash))
            for locale, title in [("zh", "项目经历"), ("en", "Projects")]:
                path = root / f"synthetic-{locale}.docx"
                preview = build_standard_resume(path, profile={"candidate_display_name": "合成候选人", "email": "synthetic@example.test", "minimum_salary": 100000}, claim_set=approved, locale=locale)
                paragraphs = [item.text for item in Document(path).paragraphs]
                self.assertEqual(paragraphs.count(wording), 1)
                self.assertIn(title, paragraphs)
                self.assertNotIn("100000", " ".join(paragraphs))
                self.assertEqual([item["category"] for item in preview["sections"]], ["project"])
            tampered = deepcopy(approved)
            tampered["claims"][0]["allowed_wording"] = ["Unapproved achievement"]
            with self.assertRaises(JobOpsError):
                build_standard_resume(root / "bad.docx", profile={"candidate_display_name": "Synthetic"}, claim_set=tampered, locale="en")
            self.assertFalse((root / "bad.docx").exists())

    @patch.dict(os.environ, {"JOBFLOW_PDF_EXPORT_PROVIDER": "word"})
    def test_capability_discovery_does_not_claim_real_conversion_or_launch_programs(self):
        with patch("jobops.document_capabilities.word_registered", return_value=True), patch("subprocess.run") as run:
            status = document_capabilities()
            self.assertEqual(status["pdf_export_provider"], "WORD")
            self.assertEqual(status["pdf_export"], "DETECTED_NOT_VERIFIED")
            run.assert_not_called()
        self.assertNotIn("path", json.dumps(status))


if __name__ == "__main__":
    unittest.main()
