from __future__ import annotations

import unittest
from unittest.mock import patch

from _support import PROJECT, project_temp
import json
from jobops.document_qa import compare_text, expected_pdf_text, reading_order_issues, repeated_story_issues, structural_qa


class DocumentReadingOrderTests(unittest.TestCase):
    def test_original_complex_release_fixture_is_preserved_and_rejected(self):
        root = PROJECT / "tests/fixtures/release-document-qa"
        result = structural_qa(root / "complex-master-resume.docx", root / "complex-master-resume.pdf",
                               sorted(root.glob("complex-master-resume-page-*.png")),
                               visual_record=json.loads((root / "visual-review.json").read_text()), page_limit=2)
        self.assertEqual(result.status, "FAIL")
        self.assertGreater(result.reading_order_issue_count, 0)
        self.assertLess(result.text_difference_ratio, 0.08)

    def test_interleaved_columns_fail_even_when_all_tokens_remain(self):
        original = "Alpha: I tested 12 cases in 2024; the team reported 20% fewer errors.\nBeta: I documented 24 cases in 2025; no financial outcome was measured."
        interleaved = "Alpha: I tested 12 cases in 2024; the team Beta: I documented 24 cases in 2025; no financial\nreported 20% fewer errors. outcome was measured."
        self.assertEqual(compare_text(original, interleaved), 0)
        self.assertEqual(reading_order_issues(original, interleaved), 2)
        with project_temp() as root:
            docx, pdf = root / "synthetic.docx", root / "synthetic.pdf"
            docx.write_bytes(b"synthetic contract input")
            pdf.write_bytes(b"synthetic contract output")
            with patch("jobops.document_qa.extract_docx_text", return_value=original), patch(
                "jobops.document_qa.extract_pdf_text", return_value=(interleaved, 1)
            ), patch("jobops.document_qa.expected_pdf_text", return_value=(original, [])), patch(
                "jobops.document_qa.validate_visual_record", return_value="PASS"
            ):
                result = structural_qa(docx, pdf, [root / "page.png"], visual_record={}, page_limit=2)
            self.assertEqual(result.status, "FAIL")
            self.assertEqual(result.reading_order_issue_count, 2)
            self.assertNotIn("Alpha", str(result.as_dict()))

    def test_repeated_footer_and_page_fields_are_expected_but_not_invented(self):
        from docx import Document
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        with project_temp() as root:
            document = Document()
            document.add_paragraph("The original applicant reviewed twelve reports without claiming team savings.")
            footer = document.sections[0].footer.paragraphs[0]
            footer.add_run("Fixture - page ")
            field = OxmlElement("w:fldSimple")
            field.set(qn("w:instr"), "PAGE")
            footer._p.append(field)
            path = root / "two-pages.docx"
            document.save(path)
            expected, stories = expected_pdf_text(path, 2)
            self.assertEqual(stories, ["Fixture - page 1", "Fixture - page 2"])
            self.assertEqual(repeated_story_issues(stories, expected), 0)
            self.assertEqual(repeated_story_issues(stories, expected.replace("page 2", "page 1")), 1)
            self.assertEqual(repeated_story_issues(stories, expected.replace("Fixture - page 2", "")), 1)
            self.assertIn("without claiming team savings", expected)

    def test_line_wrapping_and_whole_paragraph_order_remain_supported(self):
        original = "Alpha: I checked 12 stores in 2024 with the team.\nBeta: I documented 24 reports in 2025 without measured savings."
        wrapped = "Beta: I documented 24 reports in 2025\nwithout measured savings.\nAlpha: I checked 12 stores\nin 2024 with the team."
        self.assertEqual(reading_order_issues(original, wrapped), 0)

    def test_swapped_numbers_or_responsibility_fail_without_changing_word_counts(self):
        original = "I checked 12 stores in 2024; the team saved 20 percent in 2025."
        for changed in ("I checked 20 stores in 2024; the team saved 12 percent in 2025.",
                        "the team checked 12 stores in 2024; I saved 20 percent in 2025."):
            self.assertEqual(compare_text(original, changed), 0)
            self.assertEqual(reading_order_issues(original, changed), 1)
        self.assertEqual(reading_order_issues("2024 年核验 12 家门店，团队节约 3.5 万元。", "2024 年核验 3.5 家门店，团队节约 12 万元。"), 1)


if __name__ == "__main__":
    unittest.main()
