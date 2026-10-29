"""Build a separate single-column release sample; retain the complex negative."""
from pathlib import Path
import shutil
import sys

PROJECT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(PROJECT / "src"))

from docx import Document
from docx.shared import Inches, Pt
from jobops.document_builder import export_docx_to_pdf, render_pdf_to_pngs
from jobops.document_qa import automated_visual_probe, extract_docx_text
from jobops.util import write_json


def main():
    output = Path(__file__).resolve().parent
    source = PROJECT / "tests/fixtures/release-document-qa/complex-master-resume.docx"
    # A deliberately new layout, not a format-preservation operation. Exact
    # synthetic factual wording is retained; only the old layout footer changes.
    lines = extract_docx_text(source).splitlines()
    # The old table stores both headers before both values. Restore the known
    # synthetic cell associations explicitly, not by trusting flattened order.
    skills = lines.index("SKILLS")
    assert lines[skills + 1] == "EDUCATION"
    skill_fact, education_fact = lines[skills + 2:skills + 4]
    assert skill_fact == "Evidence analysis, Python, and SQL using synthetic inputs."
    assert education_fact == "Synthetic university fixture - not a real credential."
    lines[skills:skills + 4] = ["SKILLS", skill_fact, "EDUCATION", education_fact]
    document = Document()
    document.core_properties.author = document.core_properties.last_modified_by = ""
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.7)
    style = document.styles["Normal"]
    style.font.name, style.font.size = "Arial", Pt(11)
    for index, line in enumerate(lines):
        if line.startswith("Complex-layout format-preservation fixture"):
            continue
        if index == 0:
            document.add_heading(line, 0)
        elif line in {"PROFILE", "EXPERIENCE", "PROJECTS", "SKILLS", "EDUCATION"}:
            document.add_heading(line, 2)
        else:
            document.add_paragraph(line)
    section.footer.paragraphs[0].text = "Single-column synthetic release acceptance fixture"
    docx, pdf = output / "readable-resume.docx", output / "readable-resume.pdf"
    document.save(docx)
    export_docx_to_pdf(docx, pdf, PROJECT / ".agents/skills/job-application-operator/scripts/export-docx-pdf.ps1")
    pages = render_pdf_to_pngs(pdf, output, shutil.which("pdftoppm") or "pdftoppm")
    renamed = []
    for index, page in enumerate(pages, 1):
        target = output / f"readable-resume-page-{index}.png"
        page.replace(target)
        renamed.append(target)
    write_json(output / "automated-render-probe.json", automated_visual_probe(renamed))
    print("Rendered synthetic sample. A separate page inspection must produce visual-review.json.")


if __name__ == "__main__":
    main()
