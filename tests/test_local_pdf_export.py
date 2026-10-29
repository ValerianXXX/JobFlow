from __future__ import annotations

import os
from pathlib import Path
import subprocess
import sys
import unittest
from unittest.mock import patch
import zipfile

from _support import project_temp
from docx import Document
from jobops.document_builder import export_docx_to_pdf
from jobops.document_capabilities import document_capabilities
from jobops.errors import JobOpsError
from jobops.local_pdf_export import _conversion_workspace, _run_converter, _validate_input, export_with_libreoffice, libreoffice_binary
from workbench_fixtures import text_pdf


class LocalPdfExportTests(unittest.TestCase):
    @unittest.skipUnless(os.name == "nt", "Windows MAX_PATH cleanup regression")
    def test_long_office_cache_paths_are_removed_after_conversion(self):
        with project_temp() as parent:
            with _conversion_workspace(parent) as root:
                path = root / ("a" * 80) / ("b" * 80) / ("c" * 80)
                extended = Path("\\\\?\\" + str(path))
                extended.mkdir(parents=True)
                (extended / "synthetic-cache.txt").write_text("synthetic cache")
                self.assertGreater(len(str(path)), 260)
            self.assertFalse(root.exists())

    def test_configuration_is_explicit_and_discovery_does_not_launch(self):
        with patch.dict(os.environ, {"JOBFLOW_PDF_EXPORT_PROVIDER": "libreoffice"}), patch(
            "jobops.local_pdf_export.libreoffice_binary", return_value=Path("synthetic-soffice")
        ), patch("subprocess.run") as run:
            capability = document_capabilities()
            self.assertEqual(capability["pdf_export_provider"], "LIBREOFFICE")
            self.assertEqual(capability["pdf_export"], "DETECTED_NOT_VERIFIED")
            self.assertFalse(capability["automatic_installation"])
            self.assertNotIn("synthetic-soffice", str(capability))
            run.assert_not_called()
        with patch.dict(os.environ, {"JOBFLOW_PDF_EXPORT_PROVIDER": "mistyped-private-value"}):
            self.assertEqual(document_capabilities()["pdf_export_configuration"], "INVALID")
            with self.assertRaises(JobOpsError) as caught:
                export_docx_to_pdf(Path("input.docx"), Path("output.pdf"), Path("export.ps1"))
            self.assertEqual(caught.exception.code, "DOCX_PDF_PROVIDER_INVALID")
            self.assertNotIn("mistyped", str(caught.exception.as_dict()))

    def test_invalid_explicit_binary_never_falls_back(self):
        with patch.dict(os.environ, {"JOBFLOW_LIBREOFFICE_BINARY": "relative/soffice.com"}), patch("shutil.which") as which:
            self.assertIsNone(libreoffice_binary())
            which.assert_not_called()

    def test_selected_converter_does_not_launch_word_after_failure(self):
        with patch.dict(os.environ, {"JOBFLOW_PDF_EXPORT_PROVIDER": "libreoffice"}), patch(
            "jobops.local_pdf_export.export_with_libreoffice", side_effect=JobOpsError("DOCX_PDF_EXPORT_FAILED", "blocked")
        ), patch("jobops.document_builder.subprocess.run") as word:
            with self.assertRaises(JobOpsError):
                export_docx_to_pdf(Path("input.docx"), Path("output.pdf"), Path("export.ps1"))
            word.assert_not_called()

    def test_failed_conversion_never_accepts_stale_output_and_cleans_profile(self):
        with project_temp() as root:
            source, target = root / "source.docx", root / "old.pdf"
            Document().save(source)
            original = source.read_bytes()
            target.write_bytes(b"unchanged previous PDF")
            outcomes = [0, 1, subprocess.TimeoutExpired(["sensitive-path"], 90), OSError("private details")]
            for outcome in outcomes:
                with self.subTest(outcome=type(outcome).__name__):
                    def fail(command, *, cwd, env, timeout):
                        self.assertEqual(timeout, 90)
                        self.assertTrue(any(arg.startswith("-env:UserInstallation=") for arg in command))
                        self.assertNotIn("PRIVATE_TEST_KEY", env)
                        self.assertNotIn("--accept", " ".join(command))
                        self.assertTrue(Path(env["TEMP"]).is_relative_to(cwd))
                        if isinstance(outcome, Exception):
                            raise outcome
                        return outcome
                    with patch.dict(os.environ, {"PRIVATE_TEST_KEY": "do-not-inherit"}), patch(
                        "jobops.local_pdf_export.libreoffice_binary", return_value=root / "soffice.com"
                    ), patch("jobops.local_pdf_export._run_converter", side_effect=fail):
                        with self.assertRaises(JobOpsError) as caught:
                            export_with_libreoffice(source, target)
                    self.assertNotIn("private details", str(caught.exception.as_dict()))
                    self.assertNotIn("sensitive-path", str(caught.exception.as_dict()))
                    self.assertEqual(target.read_bytes(), b"unchanged previous PDF")
                    self.assertEqual(source.read_bytes(), original)
                    self.assertEqual(list(root.glob("jobflow-pdf-*")), [])

    def test_fresh_pdf_is_committed_and_work_directory_is_unique(self):
        with project_temp() as root:
            source, target = root / "source.docx", root / "out.pdf"
            Document().save(source)
            pdf = text_pdf([(48, 740, "SYNTHETIC 2024 20%")])
            workspaces = []
            def convert(command, *, cwd, env, timeout):
                # DOS aliases can be reused after deletion; compare the actual
                # newly created directory identities while each still exists.
                workspaces.append(cwd.resolve())
                self.assertEqual(Path(command[-1]).read_bytes(), source.read_bytes())
                (cwd / "output" / "input.pdf").write_bytes(pdf)
                return 0
            with patch("jobops.local_pdf_export.libreoffice_binary", return_value=root / "soffice.com"), patch(
                "jobops.local_pdf_export._run_converter", side_effect=convert
            ):
                export_with_libreoffice(source, target)
                export_with_libreoffice(source, target)
            self.assertEqual(target.read_bytes(), pdf)
            self.assertNotEqual(workspaces[0], workspaces[1])
            self.assertTrue(all(not folder.exists() for folder in workspaces))

    def test_active_content_is_rejected_before_any_process(self):
        with project_temp() as root:
            clean = root / "clean.docx"
            Document().save(clean)
            _validate_input(clean)
            rel = '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="r1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="https://untrusted.example.test/image" TargetMode="External"/></Relationships>'
            fields = '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:instrText>DDEAUTO dangerous</w:instrText></w:r></w:p></w:body></w:document>'
            cases = [("word/vbaProject.bin", b"macro"), ("word/embeddings/oleObject.bin", b"ole"),
                     ("word/_rels/document.xml.rels", rel.encode()), ("word/document.xml", fields.encode())]
            for index, (part, value) in enumerate(cases):
                bad = root / f"bad-{index}.docx"
                with zipfile.ZipFile(clean) as src, zipfile.ZipFile(bad, "w") as out:
                    for item in src.infolist():
                        if item.filename != part:
                            out.writestr(item, src.read(item))
                    out.writestr(part, value)
                with patch("jobops.local_pdf_export.libreoffice_binary", return_value=root / "soffice.com"), patch(
                    "jobops.local_pdf_export._run_converter"
                ) as run:
                    with self.assertRaises(JobOpsError) as caught:
                        export_with_libreoffice(bad, root / "out.pdf")
                    self.assertEqual(caught.exception.code, "DOCX_PDF_INPUT_UNSUPPORTED")
                    run.assert_not_called()

    def test_owned_process_does_not_leave_a_child_after_timeout(self):
        with project_temp() as root:
            marker = root / "late-child.txt"
            child_code = "import pathlib,sys,time;time.sleep(2);pathlib.Path(sys.argv[1]).write_text('escaped')"
            parent = "import subprocess,sys,time;subprocess.Popen([sys.executable,'-c',sys.argv[1],sys.argv[2]]);time.sleep(20)"
            with self.assertRaises(subprocess.TimeoutExpired):
                _run_converter([str(Path(sys.executable).resolve()), "-c", parent, child_code, str(marker)],
                               cwd=root, env=dict(os.environ), timeout=0.5)
            # This wait tests that the killed child cannot write after cleanup.
            import time
            time.sleep(2.2)
            self.assertFalse(marker.exists())

    def test_owned_process_returns_exit_without_inheriting_control_pipe(self):
        with project_temp() as root:
            self.assertEqual(_run_converter([str(Path(sys.executable).resolve()), "-c", "import sys;sys.stdin.read();sys.exit(7)"],
                                            cwd=root, env=dict(os.environ), timeout=10), 7)


if __name__ == "__main__":
    unittest.main()
